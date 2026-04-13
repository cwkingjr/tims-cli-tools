from pathlib import Path
from typing import Protocol

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill

from . import file_utils


class ConsolidatedSpreadsheetWriter(Protocol):
    def write(
        self,
        *,
        original_df: pd.DataFrame,
        all_payments_df: pd.DataFrame,
        all_payments_totals_df: pd.DataFrame,
        payee_tuples: list[tuple[str, pd.DataFrame]],
    ) -> Path: ...


class IndividualSpreadsheetWriter(Protocol):
    def write(
        self,
        *,
        pay_to_name: str,
        individual_payments_totals_df: pd.DataFrame,
        individual_payments_df: pd.DataFrame,
    ) -> Path: ...


class XlsxConsolidatedSpreadsheetWriter:
    def write(
        self,
        *,
        original_df: pd.DataFrame,
        all_payments_df: pd.DataFrame,
        all_payments_totals_df: pd.DataFrame,
        payee_tuples: list[tuple[str, pd.DataFrame]],
    ) -> Path:
        file_path = file_utils.create_new_spreadsheet_central_tz_filepath(
            filename_prefix="payroll_consolidated"
        )

        underlined_payees = []

        with pd.ExcelWriter(file_path, engine="xlsxwriter") as writer:
            original_df.to_excel(writer, sheet_name="Original_Data", index=False)
            all_payments_totals_df.to_excel(
                writer, sheet_name="All_Payments_Totals", index=False
            )
            all_payments_df.to_excel(writer, sheet_name="All_Payments", index=False)
            for one_payee_tuple in payee_tuples:
                payee, df = one_payee_tuple
                underlined_payee = payee.replace(" ", "_")
                underlined_payees.append(underlined_payee)
                df.to_excel(writer, sheet_name=underlined_payee, index=False)

        self._format(file_path, underlined_payees)
        return file_path

    def _format(self, file_path: Path, payees: list[str]) -> None:
        currency_format = "#,##0.00"
        header_fill = PatternFill(start_color="FF03BDFC", fill_type="solid")

        workbook = load_workbook(file_path)

        self._format_original_data(
            workbook["Original_Data"], header_fill, currency_format
        )
        self._format_payments_totals(
            workbook["All_Payments_Totals"], header_fill, currency_format
        )
        self._format_payments(workbook["All_Payments"], header_fill, currency_format)

        for payee in payees:
            self._format_payments(workbook[payee], header_fill, currency_format)

        workbook.save(file_path)

    def _format_original_data(
        self, sheet, header_fill: PatternFill, currency_format: str
    ) -> None:
        for cell in sheet[1]:
            cell.fill = header_fill
        for column_letter in "BCDEFGHIJKLMNOPQ":
            sheet.column_dimensions[column_letter].width = 15
            for cell in sheet[column_letter]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in sheet["B"]:
            cell.number_format = "yyyy-mm-dd"

    def _format_payments_totals(
        self, sheet, header_fill: PatternFill, currency_format: str
    ) -> None:
        for cell in sheet[1]:
            cell.fill = header_fill
        for column_letter in "ABCDEFGHIJK":
            sheet.column_dimensions[column_letter].width = 15
            for cell in sheet[column_letter]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_letter in "BCDEFGHIJ":
            for cell in sheet[column_letter]:
                cell.number_format = currency_format

    def _format_payments(
        self, sheet, header_fill: PatternFill, currency_format: str
    ) -> None:
        for cell in sheet[1]:
            cell.fill = header_fill
        for column_letter in "GHIJKLMNO":
            for cell in sheet[column_letter]:
                cell.number_format = currency_format
        for column_letter in "ABCDEFGHIJKLMNOPQ":
            sheet.column_dimensions[column_letter].width = 15
            for cell in sheet[column_letter]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in sheet["P"]:
            cell.number_format = "hh:mm"
        for cell in sheet["C"]:
            cell.number_format = "yyyy-mm-dd"


class XlsxIndividualSpreadsheetWriter:
    def write(
        self,
        *,
        pay_to_name: str,
        individual_payments_totals_df: pd.DataFrame,
        individual_payments_df: pd.DataFrame,
    ) -> Path:
        clean_payto = pay_to_name.lower().replace(" ", "_")
        file_path = file_utils.create_new_spreadsheet_central_tz_filepath(
            filename_prefix=f"payroll_{clean_payto}"
        )

        with pd.ExcelWriter(file_path, engine="xlsxwriter") as writer:
            individual_payments_totals_df.to_excel(
                writer, sheet_name="Payment_Totals", index=False
            )
            individual_payments_df.to_excel(writer, sheet_name="Payments", index=False)

        self._format(file_path)
        return file_path

    def _format(self, file_path: Path) -> None:
        currency_format = "#,##0.00"
        header_fill = PatternFill(start_color="FF03BDFC", fill_type="solid")

        workbook = load_workbook(file_path)

        payments = workbook["Payments"]
        for cell in payments[1]:
            cell.fill = header_fill
        for column_letter in "GHIJKLMNO":
            for cell in payments[column_letter]:
                cell.number_format = currency_format
        for column_letter in "ABCDEFGHIJKLMNOPQ":
            payments.column_dimensions[column_letter].width = 15
            for cell in payments[column_letter]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in payments["P"]:
            cell.number_format = "hh:mm"
        for cell in payments["C"]:
            cell.number_format = "yyyy-mm-dd"

        payment_totals = workbook["Payment_Totals"]
        for cell in payment_totals[1]:
            cell.fill = header_fill
        for column_letter in "ABCDEFGHIJK":
            payment_totals.column_dimensions[column_letter].width = 15
            for cell in payment_totals[column_letter]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_letter in "BCDEFGHIJ":
            for cell in payment_totals[column_letter]:
                cell.number_format = currency_format

        workbook.save(file_path)
