from datetime import date, time

import pandas as pd
import pytest
from openpyxl import load_workbook

from tims_cli_tools import field
from tims_cli_tools.payroll_writers import (
    XlsxConsolidatedSpreadsheetWriter,
    XlsxIndividualSpreadsheetWriter,
)


class TestXlsxConsolidatedSpreadsheetWriterIntegration:
    """Integration tests for the consolidated spreadsheet writer."""

    @pytest.fixture
    def sample_data(self):
        original_df = pd.DataFrame(
            {
                field.BU: [12345],
                field.STRUCTURE: [3],
                field.DATE: [date(2025, 1, 15)],
            }
        )

        all_payments_df = pd.DataFrame(
            {
                "bu": [12345],
                "inspection_date": [date(2025, 1, 15)],
                "crew_lead": ["John Doe"],
                "pay_to": ["John Doe"],
                "structure_type": ["Guyed"],
                "tia_inspection": [5.0],
                "additional_canister_level": [0.0],
                "hvf": [10.0],
                "lighting_inspection_price": [5.0],
                "migratory_bird": [2.5],
                "windsim": [4.0],
                "tension": [10.0],
                "hr_pay": [15.0],
                "site_total": [51.5],
                "maintenance": [time(hour=0, minute=30)],
                "extra_cans": [2],
            }
        )

        all_payments_totals_df = pd.DataFrame(
            {
                "pay_to": ["John Doe"],
                "tia_inspection": [5.0],
                "additional_canister_level": [0.0],
                "hvf": [10.0],
                "lighting_inspection_price": [5.0],
                "migratory_bird": [2.5],
                "windsim": [4.0],
                "tension": [10.0],
                "hr_pay": [15.0],
                "site_total": [51.5],
                "extra_cans": [2],
            }
        )

        payee_df = pd.DataFrame(
            {
                "bu": [12345],
                "pay_to": ["John Doe"],
                "site_total": [51.5],
            }
        )

        payee_tuples = [("John Doe", payee_df)]

        return original_df, all_payments_df, all_payments_totals_df, payee_tuples

    def test_creates_all_expected_sheets(self, sample_data, tmp_path):
        original_df, all_payments_df, all_payments_totals_df, payee_tuples = sample_data

        writer = XlsxConsolidatedSpreadsheetWriter()
        output_path = writer.write(
            original_df=original_df,
            all_payments_df=all_payments_df,
            all_payments_totals_df=all_payments_totals_df,
            payee_tuples=payee_tuples,
        )

        try:
            wb = load_workbook(output_path)
            assert "Original_Data" in wb.sheetnames
            assert "All_Payments_Totals" in wb.sheetnames
            assert "All_Payments" in wb.sheetnames
            assert "John_Doe" in wb.sheetnames
        finally:
            if output_path.exists():
                output_path.unlink()

    def test_applies_currency_formatting(self, sample_data, tmp_path):
        original_df, all_payments_df, all_payments_totals_df, payee_tuples = sample_data

        writer = XlsxConsolidatedSpreadsheetWriter()
        output_path = writer.write(
            original_df=original_df,
            all_payments_df=all_payments_df,
            all_payments_totals_df=all_payments_totals_df,
            payee_tuples=payee_tuples,
        )

        try:
            wb = load_workbook(output_path)
            sheet = wb["All_Payments_Totals"]
            cell = sheet["B2"]
            assert cell.number_format == "#,##0.00"
        finally:
            if output_path.exists():
                output_path.unlink()

    def test_applies_column_widths(self, sample_data, tmp_path):
        original_df, all_payments_df, all_payments_totals_df, payee_tuples = sample_data

        writer = XlsxConsolidatedSpreadsheetWriter()
        output_path = writer.write(
            original_df=original_df,
            all_payments_df=all_payments_df,
            all_payments_totals_df=all_payments_totals_df,
            payee_tuples=payee_tuples,
        )

        try:
            wb = load_workbook(output_path)
            sheet = wb["All_Payments"]
            assert sheet.column_dimensions["A"].width > 0
        finally:
            if output_path.exists():
                output_path.unlink()


class TestXlsxIndividualSpreadsheetWriterIntegration:
    """Integration tests for the individual spreadsheet writer."""

    @pytest.fixture
    def sample_data(self):
        individual_payments_totals_df = pd.DataFrame(
            {
                "pay_to": ["John Doe"],
                "tia_inspection": [5.0],
                "additional_canister_level": [0.0],
                "hvf": [10.0],
                "lighting_inspection_price": [5.0],
                "migratory_bird": [2.5],
                "windsim": [4.0],
                "tension": [10.0],
                "hr_pay": [15.0],
                "site_total": [51.5],
                "extra_cans": [2],
            }
        )

        individual_payments_df = pd.DataFrame(
            {
                "bu": [12345],
                "pay_to": ["John Doe"],
                "site_total": [51.5],
            }
        )

        return individual_payments_totals_df, individual_payments_df

    def test_creates_expected_sheets(self, sample_data, tmp_path):
        totals_df, payments_df = sample_data

        writer = XlsxIndividualSpreadsheetWriter()
        output_path = writer.write(
            pay_to_name="John Doe",
            individual_payments_totals_df=totals_df,
            individual_payments_df=payments_df,
        )

        try:
            wb = load_workbook(output_path)
            assert "Payment_Totals" in wb.sheetnames
            assert "Payments" in wb.sheetnames
        finally:
            if output_path.exists():
                output_path.unlink()

    def test_applies_currency_formatting_to_payments(self, sample_data, tmp_path):
        totals_df, payments_df = sample_data

        writer = XlsxIndividualSpreadsheetWriter()
        output_path = writer.write(
            pay_to_name="John Doe",
            individual_payments_totals_df=totals_df,
            individual_payments_df=payments_df,
        )

        try:
            wb = load_workbook(output_path)
            sheet = wb["Payments"]
            cell = sheet["G2"]
            assert cell.number_format == "#,##0.00"
        finally:
            if output_path.exists():
                output_path.unlink()

    def test_applies_column_widths(self, sample_data, tmp_path):
        totals_df, payments_df = sample_data

        writer = XlsxIndividualSpreadsheetWriter()
        output_path = writer.write(
            pay_to_name="John Doe",
            individual_payments_totals_df=totals_df,
            individual_payments_df=payments_df,
        )

        try:
            wb = load_workbook(output_path)
            sheet = wb["Payments"]
            assert sheet.column_dimensions["A"].width > 0
        finally:
            if output_path.exists():
                output_path.unlink()
